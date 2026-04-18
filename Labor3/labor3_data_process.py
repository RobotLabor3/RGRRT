import pandas as pd
import numpy as np
import os
from openpyxl import load_workbook
from openpyxl.styles import Font


folder_path = ""
file_path = folder_path + "plan_result.txt"

output_dir = os.path.dirname(file_path)
output_file = os.path.join(output_dir, "planner_result.xlsx")

data = {}

with open(file_path, "r") as f:
    lines = [line.strip() for line in f.readlines()]

group_id = 1
i = 0

while i < len(lines):
    line = lines[i]

    if "-----" in line:
        group_id += 1
        i += 1
        continue

    if "plan_time" in line:
        planner = line.split()[0]

        if group_id not in data:
            data[group_id] = {}
        if planner not in data[group_id]:
            data[group_id][planner] = {"time": [], "length": []}

        values = []
        j = i + 1

        while j < len(lines):
            next_line = lines[j]
            if "plan_time" in next_line or "-----" in next_line:
                break

            for p in next_line.split():
                try:
                    values.append(float(p))
                except:
                    pass

            j += 1

        if len(values) >= 2:
            n = len(values) // 2
            times = values[:n]
            lengths = values[n:n*2]

            for t, l in zip(times, lengths):
                data[group_id][planner]["time"].append(t)
                data[group_id][planner]["length"].append(l)

        i = j
        continue

    i += 1


def stats(arr):
    if len(arr) == 0:
        return [np.nan]*4
    return [
        np.mean(arr),
        np.max(arr),
        np.min(arr),
        np.median(arr)
    ]


rows = []

for g, planners in data.items():
    for planner, vals in planners.items():
        t_stats = stats(vals["time"])
        l_stats = stats(vals["length"])

        rows.append({
            "group": g,
            "planner": planner,

            "success_count": len(vals["length"]),

            "time_mean": t_stats[0],
            "time_max": t_stats[1],
            "time_min": t_stats[2],
            "time_median": t_stats[3],

            "length_mean": l_stats[0],
            "length_max": l_stats[1],
            "length_min": l_stats[2],
            "length_median": l_stats[3],
        })

df = pd.DataFrame(rows).sort_values(by=["group", "planner"])

df.to_excel(output_file, index=False)

wb = load_workbook(output_file)
ws = wb.active

bold_font = Font(bold=True)

headers = [cell.value for cell in ws[1]]

group_col = headers.index("group") + 1
mean_col = headers.index("length_mean") + 1
median_col = headers.index("length_median") + 1

group_rows = {}

for row in range(2, ws.max_row + 1):
    g = ws.cell(row=row, column=group_col).value
    mean_val = ws.cell(row=row, column=mean_col).value
    median_val = ws.cell(row=row, column=median_col).value

    group_rows.setdefault(g, []).append((row, mean_val, median_val))


for g, items in group_rows.items():
    valid = [(r, m) for r, m, _ in items if m is not None]
    if valid:
        min_mean = min(m for _, m in valid)
        for r, m in valid:
            if m == min_mean:
                ws.cell(row=r, column=mean_col).font = bold_font


for g, items in group_rows.items():
    valid = [(r, md) for r, _, md in items if md is not None]
    if valid:
        min_median = min(md for _, md in valid)
        for r, md in valid:
            if md == min_median:
                ws.cell(row=r, column=median_col).font = bold_font


wb.save(output_file)
